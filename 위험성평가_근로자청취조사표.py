import io
import uuid
from datetime import datetime

import pandas as pd
import streamlit as st
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH

st.set_page_config(page_title="위험성평가 근로자청취조사표", layout="wide")

if "risk_interview_records" not in st.session_state:
    st.session_state.risk_interview_records = []
if "last_submission_id" not in st.session_state:
    st.session_state.last_submission_id = None


def build_response_docx(record):
    """근로자청취조사 응답 1건을 .docx 보고서로 생성"""
    doc = Document()

    title = doc.add_heading("위험성평가 근로자 청취조사표", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_heading("1. 기본 정보", level=1)
    basic_table = doc.add_table(rows=0, cols=2)
    basic_table.style = "Table Grid"
    for label, value in [
        ("소속", record["소속"]),
        ("부서", record["부서"]),
        ("성명", record["성명"]),
        ("담당 업무", record["담당업무"]),
    ]:
        row = basic_table.add_row().cells
        row[0].text, row[1].text = label, str(value)

    doc.add_heading("2. 산업재해 / 아차사고 경험", level=1)
    doc.add_paragraph(f"산업재해 혹은 사고 경험: {record['산업재해_사고경험']}")
    doc.add_paragraph(f"아차사고 경험: {record['아차사고경험']}")
    doc.add_paragraph(f"발생경위: {record['발생경위']}")

    doc.add_heading("3. 유해·위험요인 및 개선 의견", level=1)
    doc.add_paragraph(record["개선의견"])

    footer = doc.add_paragraph(f"제출일시: {record['제출일시']}")
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


st.markdown(
    """
    <style>
    .stat-card {
        border: 1px solid rgba(11,11,11,0.10);
        border-radius: 12px;
        padding: 20px 24px;
        background: #fcfcfb;
    }
    .stat-card .stat-label {
        font-family: system-ui, -apple-system, "Segoe UI", sans-serif;
        font-size: 15px;
        font-weight: 600;
        color: #52514e;
        margin-bottom: 6px;
    }
    .stat-card .stat-value {
        font-family: system-ui, -apple-system, "Segoe UI", sans-serif;
        font-size: 44px;
        font-weight: 800;
        color: #0b0b0b;
        line-height: 1.1;
    }
    .stat-card .stat-sub {
        font-family: system-ui, -apple-system, "Segoe UI", sans-serif;
        font-size: 13px;
        color: #898781;
        margin-top: 4px;
    }
    .stat-card.accent-total { border-left: 6px solid #2a78d6; }
    .stat-card.accent-accident { border-left: 6px solid #e2574c; }
    .stat-card.accent-nearmiss { border-left: 6px solid #fab219; }
    @media (prefers-color-scheme: dark) {
        .stat-card { background: #1a1a19; border-color: rgba(255,255,255,0.10); }
        .stat-card .stat-label { color: #c3c2b7; }
        .stat-card .stat-value { color: #ffffff; }
        .stat-card .stat-sub { color: #898781; }
    }
    </style>
    """,
    unsafe_allow_html=True,
)

st.title("2026년 위험성평가 근로자 청취조사표")
st.caption(
    "본 설문조사는 『2026년도 위험성평가』 실시 전 작업환경의 유해·위험요인을 파악하고 "
    "종사자의 의견을 청취할 목적으로 실시합니다. 여러분들의 의견을 통해 더 안전하고 "
    "건강한 일터를 실현할 수 있으니 많은 협조바랍니다."
)

tab1, tab2 = st.tabs(["설문 응답 입력", "응답 결과 관리"])

# =====================================================================
# TAB 1. 설문 응답 입력 (근로자용)
# =====================================================================
with tab1:
    with st.form("risk_interview_form", clear_on_submit=True):
        st.markdown("#### 1. 소속, 부서, 성명은 어떻게 됩니까?")
        col1, col2, col3 = st.columns(3)
        with col1:
            affiliation = st.text_input("소속 *", placeholder="예: 서울병원")
        with col2:
            department = st.text_input("부서 *")
        with col3:
            respondent_name = st.text_input("성명 *")

        st.markdown("#### 2. 하시는 업무는 무엇입니까?")
        job_duty = st.text_input("담당 업무 *", key="job_duty")

        st.markdown("#### 3. 업무와 관련하여 일하시는 중 산업재해 혹은 사고 경험이 있습니까?")
        accident_exp = st.radio(
            "산업재해/사고 경험", ["예", "아니오"], horizontal=True,
            label_visibility="collapsed", key="accident_exp",
        )

        st.markdown("#### 4. 업무와 관련하여 일하시는 중 아차사고 경험이 있습니까?")
        st.caption(
            "*아차사고: 사고가 발생할 뻔하였으나 직접적인 사고로 이어지지 않은 상황으로, "
            "다치지는 않았으나 부상으로 이어질 가능성이 있었던 경험"
        )
        near_miss_exp = st.radio(
            "아차사고 경험", ["예", "아니오"], horizontal=True,
            label_visibility="collapsed", key="near_miss_exp",
        )

        st.markdown("#### 5. (3번 혹은 4번에 '예'로 답하신 경우) 업무 중 산업재해 혹은 (아차)사고는 어떠한 경위로 발생하였습니까?")
        st.caption("※ 6하 원칙(누가, 언제, 어디서, 어떻게, 왜)으로 작성해주시면 감사하겠습니다.")
        incident_detail = st.text_area("발생경위", height=100, label_visibility="collapsed")

        st.markdown(
            "#### 6. 작업장소 혹은 업무 상 유해·위험요인으로 생각되거나 "
            "안전한 작업환경을 위해 개선해야 할 사항을 자유롭게 작성바랍니다."
        )
        improvement_opinion = st.text_area("개선 의견", height=120, label_visibility="collapsed")

        submitted = st.form_submit_button("설문 제출", use_container_width=True)

        if submitted:
            job_duty_value = job_duty.strip()
            if not affiliation or not department or not respondent_name:
                st.error("소속, 부서, 성명은 필수 입력 항목입니다.")
            elif not job_duty_value:
                st.error("담당 업무를 작성해주세요.")
            elif (accident_exp == "예" or near_miss_exp == "예") and not incident_detail.strip():
                st.error("3번 또는 4번에 '예'로 응답하신 경우, 5번 발생경위를 작성해주세요.")
            else:
                record = {
                    "id": uuid.uuid4().hex[:8],
                    "소속": affiliation,
                    "부서": department,
                    "성명": respondent_name,
                    "담당업무": job_duty_value,
                    "산업재해_사고경험": accident_exp,
                    "아차사고경험": near_miss_exp,
                    "발생경위": incident_detail.strip() or "-",
                    "개선의견": improvement_opinion.strip() or "-",
                    "제출일시": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
                }
                st.session_state.risk_interview_records.append(record)
                st.session_state.last_submission_id = record["id"]
                st.success(f"{respondent_name}님의 설문이 제출되었습니다. 협조해 주셔서 감사합니다.")

    last_id = st.session_state.last_submission_id
    if last_id:
        last_record = next(
            (r for r in st.session_state.risk_interview_records if r["id"] == last_id), None
        )
        if last_record:
            st.markdown("#### 제출한 내용 파일로 저장")
            st.download_button(
                "📄 워드(.docx)로 저장",
                data=build_response_docx(last_record),
                file_name=f"위험성평가_근로자청취조사표_{last_record['성명']}.docx",
                mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                use_container_width=True,
                key=f"download_docx_{last_id}",
            )

# =====================================================================
# TAB 2. 응답 결과 관리 (관리자용)
# =====================================================================
with tab2:
    st.subheader("제출된 응답 현황")

    records = st.session_state.risk_interview_records
    all_affiliations = sorted({r["소속"] for r in records})
    filter_col, _ = st.columns([1, 3])
    with filter_col:
        affiliation_filter = st.selectbox("소속 필터", ["전체"] + all_affiliations, key="affiliation_filter")

    filtered = [r for r in records if affiliation_filter == "전체" or r["소속"] == affiliation_filter]
    total_count = len(filtered)
    accident_count = sum(1 for r in filtered if r["산업재해_사고경험"] == "예")
    near_miss_count = sum(1 for r in filtered if r["아차사고경험"] == "예")

    stat_col1, stat_col2, stat_col3 = st.columns(3)
    with stat_col1:
        st.markdown(
            f"""
            <div class="stat-card accent-total">
                <div class="stat-label">총 응답 건수{'' if affiliation_filter == '전체' else f' ({affiliation_filter})'}</div>
                <div class="stat-value">{total_count}건</div>
            </div>
            """,
            unsafe_allow_html=True,
        )
    with stat_col2:
        st.markdown(
            f"""
            <div class="stat-card accent-accident">
                <div class="stat-label">⚠ 산업재해/사고 경험</div>
                <div class="stat-value">{accident_count}건</div>
            </div>
            """,
            unsafe_allow_html=True,
        )
    with stat_col3:
        st.markdown(
            f"""
            <div class="stat-card accent-nearmiss">
                <div class="stat-label">◆ 아차사고 경험</div>
                <div class="stat-value">{near_miss_count}건</div>
            </div>
            """,
            unsafe_allow_html=True,
        )

    st.divider()
    st.caption(f"현재 필터: {affiliation_filter}  ·  표시 건수: {total_count}건")

    if filtered:
        df = pd.DataFrame(filtered).drop(columns=["id"])
        st.dataframe(df, use_container_width=True)

        csv = df.to_csv(index=False).encode("utf-8-sig")
        st.download_button(
            "CSV로 다운로드",
            data=csv,
            file_name="위험성평가_근로자청취조사_결과.csv",
            mime="text/csv",
            use_container_width=True,
        )

        if st.button("전체 응답 초기화"):
            st.session_state.risk_interview_records = []
            st.session_state.last_submission_id = None
            st.rerun()
    else:
        if records:
            st.info(f"'{affiliation_filter}' 소속에 해당하는 응답이 없습니다.")
        else:
            st.info("아직 제출된 응답이 없습니다. [설문 응답 입력] 탭에서 설문을 작성해주세요.")
